"""Build the Word templates (.dotx) for each brand: the first set (letterhead, report, proposal, press release and
meeting notes, hand-written below) and every document in catalog.py, rendered from the blocks in word_content.py.

Light grounds for reading (the brand rule for documents), with a Dark Brown photo cover on reports and proposals.
Real heading styles (navigation pane and screen readers), en-GB language, alt text on logos, decorative art behind
text, colour pairs from the tokens.

python3 word_build.py [Brand ...]          build every template for the brands given (default: all)
python3 word_build.py --only Memo ...     build only the named templates
python3 word_build.py --copy              print every string a reader sees, for check_copy.py
"""
import copy
import os
import shutil
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor, Emu

import brand as B
import catalog as K
import ooxml_theme as T
import copy_io as C
import word_content as WC

ART = os.path.join(B.OUT, "_art")
rgb = RGBColor.from_string


# ---------- low level helpers ----------

def set_font(style_or_run, name, size=None, bold=None, italic=None, colour=None, caps=None, spacing=None):
    f = style_or_run.font
    f.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else None
    if rpr is not None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), name)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if colour:
        f.color.rgb = rgb(colour)
    if caps is not None:
        f.all_caps = caps
    if spacing is not None and rpr is not None:
        sp = rpr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            rpr.append(sp)
        sp.set(qn("w:val"), str(int(spacing * 20)))


def para_fmt(style, before=0, after=8, line=1.35, keep_next=None, left=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if left is not None:
        pf.left_indent = Cm(left)


def border(p_or_style_el, side="left", colour=B.ORANGE, size=36, space=12):
    ppr = p_or_style_el.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    el = OxmlElement("w:%s" % side)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), colour)
    bdr.append(el)


def shade(cell_or_ppr, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    cell_or_ppr.append(sh)


def anchor_picture(run, path, width_cm, x_cm, y_cm, behind=True, rel_h="page", rel_v="page", descr="", decorative=False,
                   height_cm=None):
    """Add a picture to a run and turn it into a floating picture at an absolute page position."""
    if height_cm:
        shape = run.add_picture(path, width=Cm(width_cm), height=Cm(height_cm))
    else:
        shape = run.add_picture(path, width=Cm(width_cm))
    inline = shape._inline
    ext = inline.find(qn("wp:extent"))
    docpr = inline.find(qn("wp:docPr"))
    docpr.set("descr", descr)
    if decorative:
        docpr.append(parse_xml(
            '<a:extLst %s><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
            '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/>'
            '</a:ext></a:extLst>' % nsdecls("a")))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="%d" behindDoc="%d" '
        'locked="1" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="%s"><wp:posOffset>%d</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="%s"><wp:posOffset>%d</wp:posOffset></wp:positionV>'
        '<wp:extent cx="%s" cy="%s"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>'
        % (nsdecls("wp"), 100 if behind else 251660000, 1 if behind else 0, rel_h, Cm(x_cm), rel_v, Cm(y_cm),
           ext.get("cx"), ext.get("cy")))
    anchor.append(docpr)
    anchor.append(parse_xml('<wp:cNvGraphicFramePr %s/>' % nsdecls("wp")))
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def field(run, instr):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " %s " % text
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def rich(p, text, base=None, orange=B.ORANGE, size=None, font=None):
    """{word} in orange, _word_ in italic."""
    import re
    for tok in re.split(r"(\{[^}]+\}|_[^_]+_)", text):
        if not tok:
            continue
        r = p.add_run()
        if tok.startswith("{"):
            r.text = tok[1:-1]
            r.font.color.rgb = rgb(orange)
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            r.text = tok[1:-1]
            r.italic = True
        else:
            r.text = tok
        if base:
            r.font.color.rgb = rgb(base) if not tok.startswith("{") else rgb(orange)
        if size:
            r.font.size = Pt(size)
        if font:
            set_font(r, font)
    return p


def content_control(p, alias, text, style=None):
    """Wrap a paragraph's runs in a plain-text content control so people fill the slot, not retype the layout."""
    sdt = parse_xml('<w:sdt %s><w:sdtPr><w:alias w:val="%s"/><w:tag w:val="%s"/><w:text/></w:sdtPr><w:sdtContent/></w:sdt>'
                    % (nsdecls("w"), alias, alias.lower().replace(" ", "-")))
    content = sdt.find(qn("w:sdtContent"))
    for r in list(p._p.findall(qn("w:r"))):
        content.append(r)
    p._p.append(sdt)
    return p


# ---------- styles ----------

def header_row(tbl, row=0):
    """Repeat the table style's header look as direct formatting, for apps that skip conditional table styles."""
    for cell in tbl.rows[row].cells:
        shade(cell._tc.get_or_add_tcPr(), B.DARK)
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, B.DISPLAY, 8.5, bold=True, colour=B.WHITE, caps=True)
    trpr = tbl.rows[row]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))


def table_style_xml():
    return ('<w:style %s w:type="table" w:customStyle="1" w:styleId="ClimbTable"><w:name w:val="Climb Table"/>'
            '<w:basedOn w:val="TableNormal"/><w:uiPriority w:val="40"/><w:qFormat/>'
            '<w:pPr><w:spacing w:before="60" w:after="60" w:line="276" w:lineRule="auto"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="Open Sans" w:hAnsi="Open Sans" w:cs="Open Sans"/><w:color w:val="%s"/><w:sz w:val="19"/></w:rPr>'
            '<w:tblPr><w:tblStyleRowBandSize w:val="1"/><w:tblBorders><w:bottom w:val="single" w:sz="8" w:space="0" w:color="%s"/>'
            '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/></w:tblBorders>'
            '<w:tblCellMar><w:top w:w="85" w:type="dxa"/><w:left w:w="140" w:type="dxa"/><w:bottom w:w="85" w:type="dxa"/>'
            '<w:right w:w="140" w:type="dxa"/></w:tblCellMar></w:tblPr>'
            '<w:tblStylePr w:type="firstRow"><w:rPr><w:rFonts w:ascii="Montserrat" w:hAnsi="Montserrat"/><w:b/><w:caps/>'
            '<w:color w:val="FFFFFF"/><w:sz w:val="17"/></w:rPr><w:tblPr/><w:trPr><w:tblHeader/></w:trPr>'
            '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="%s"/></w:tcPr></w:tblStylePr>'
            '<w:tblStylePr w:type="band2Horz"><w:tblPr/><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="%s"/></w:tcPr></w:tblStylePr>'
            '<w:tblStylePr w:type="lastRow"><w:rPr><w:b/><w:color w:val="%s"/></w:rPr><w:tblPr/><w:tcPr>'
            '<w:tcBorders><w:top w:val="single" w:sz="12" w:space="0" w:color="%s"/></w:tcBorders></w:tcPr></w:tblStylePr>'
            '</w:style>' % (nsdecls("w"), B.SLATE, B.CHARCOAL, B.GREY, B.DARK, B.GREY, B.CHARCOAL, B.ORANGE))


def setup_styles(doc):
    st = doc.styles
    # Document defaults: en-GB.
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = rpr_default.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr_default.append(lang)
    lang.set(qn("w:val"), "en-GB")

    n = st["Normal"]
    set_font(n, B.BODY, 10.5, colour=B.SLATE)
    para_fmt(n, 0, 8, 1.4)

    def mk(name, base="Normal", typ=WD_STYLE_TYPE.PARAGRAPH):
        try:
            return st[name]
        except KeyError:
            s = st.add_style(name, typ)
            s.base_style = st[base]
            s.quick_style = True
            return s

    t = mk("Title")
    set_font(t, B.DISPLAY_BLACK, 34, colour=B.CHARCOAL, caps=True, spacing=-0.4)
    para_fmt(t, 0, 10, 0.95)
    t.element.get_or_add_pPr()
    sub = mk("Subtitle")
    set_font(sub, B.BODY, 14, colour=B.SLATE, italic=False)
    para_fmt(sub, 0, 18, 1.3)
    for lvl, (font, size, colour, caps, before, after) in {
        1: (B.DISPLAY_XB, 22, B.CHARCOAL, True, 26, 8),
        2: (B.DISPLAY, 15, B.CHARCOAL, False, 18, 6),
        3: (B.BODY, 11.5, B.ORANGE_TEXT, False, 14, 4),
    }.items():
        h = st["Heading %d" % lvl]
        set_font(h, font, size, bold=True if lvl > 1 else False, italic=False, colour=colour, caps=caps,
                 spacing=-0.2 if lvl == 1 else 0)
        para_fmt(h, before, after, 1.1, keep_next=True)
    eb = mk("Eyebrow")
    set_font(eb, B.BODY, 9, bold=True, colour=B.ORANGE_TEXT, caps=True, spacing=1.5)
    para_fmt(eb, 12, 4, 1.2, keep_next=True)
    lead = mk("Lead")
    set_font(lead, B.BODY, 13, colour=B.SLATE)
    para_fmt(lead, 0, 12, 1.45)
    q = st["Quote"] if "Quote" in [s.name for s in st] else mk("Quote")
    set_font(q, B.DISPLAY, 15, bold=True, italic=False, colour=B.CHARCOAL)
    para_fmt(q, 14, 14, 1.3, left=0.4)
    border(q.element, "left", B.ORANGE, 36, 14)
    qs = mk("Quote Source")
    set_font(qs, B.BODY, 9.5, italic=True, colour=B.SLATE)
    para_fmt(qs, 0, 14, 1.3, left=0.4)
    callout = mk("Callout")
    set_font(callout, B.BODY, 10.5, colour=B.SLATE_GREY)
    para_fmt(callout, 10, 10, 1.4, left=0.3)
    ppr = callout.element.get_or_add_pPr()
    shade(ppr, B.GREY)
    border(callout.element, "left", B.ORANGE, 36, 10)
    big = mk("Big Number")
    set_font(big, B.DISPLAY_BLACK, 40, colour=B.ORANGE_TEXT)
    para_fmt(big, 6, 0, 1.0, keep_next=True)
    cap = mk("Caption") if "Caption" not in [s.name for s in st] else st["Caption"]
    set_font(cap, B.BODY, 9, bold=False, italic=False, colour=B.SLATE)
    para_fmt(cap, 4, 12, 1.3)
    small = mk("Small Print")
    set_font(small, B.BODY, 8, colour=B.SLATE)
    para_fmt(small, 0, 4, 1.3)
    for name in ("List Bullet", "List Number"):
        s = st[name]
        set_font(s, B.BODY, 10.5, colour=B.SLATE)
        para_fmt(s, 0, 4, 1.4)
    for name in ("Header", "Footer", "Title"):
        ppr = st[name].element.find(qn("w:pPr"))
        if ppr is not None:
            for tag in ("w:tabs", "w:pBdr"):
                el = ppr.find(qn(tag))
                if el is not None:
                    ppr.remove(el)
    hf = st["Header"]
    set_font(hf, B.BODY, 8.5, colour=B.SLATE)
    ft = st["Footer"]
    set_font(ft, B.BODY, 8, colour=B.SLATE)
    link = mk("Hyperlink", "Default Paragraph Font", WD_STYLE_TYPE.CHARACTER)
    link.font.color.rgb = rgb(B.ORANGE_DEEP)
    link.font.underline = True
    em = mk("Orange Word", "Default Paragraph Font", WD_STYLE_TYPE.CHARACTER)
    em.font.color.rgb = rgb(B.ORANGE_TEXT)
    st.element.append(parse_xml(table_style_xml()))
    # Bullet colour: Climb Orange squares via the List Bullet numbering is left to Word's default glyph for
    # compatibility; the paragraph text stays Body Slate.


def apply_theme(path, b):
    """Swap the theme part for the brand theme (colours in Word's picker, Montserrat and Open Sans as theme fonts)."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/theme/theme1.xml":
                data = T.theme_xml(b).encode()
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml")
            zout.writestr(item, data)
    os.replace(tmp, path)


def to_document(path, out):
    """A .docx copy of a template for previews."""
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
            zout.writestr(item, data)


# ---------- page furniture ----------

def page_setup(sec, top=3.4, bottom=2.8):
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(top), Cm(bottom)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.0)


def standard_header(b, header, footer, label=None, landscape=False):
    """Logo top left, URL (or a document label) top right, slab corners, legal line and page number."""
    right = 25.3 if landscape else 16.6
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(right), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run()
    lg = B.logo(b, "full-colour-light", 1024)
    r.add_picture(lg, width=Cm(4.0))
    r._r.find(".//" + qn("wp:docPr")).set("descr", C.fill(WC.shared("alt.logo"), brand=b))
    r2 = hp.add_run("\t" + (label or B.BRANDS[b]["url"]))
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = rgb(B.SLATE)
    r2.font.bold = True
    r2.font.all_caps = True
    set_font(r2, B.BODY)
    anchor_picture(hp.add_run(), os.path.join(ART, "a4_slab_tl.png"), 3.6, 0, 0, descr="", decorative=True)

    fp = footer.paragraphs[0]
    fp.text = ""
    fp.style = footer.part.document.styles["Small Print"] if False else fp.style
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(right), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run(B.LEGAL)
    r.font.size = Pt(7)
    r.font.color.rgb = rgb(B.SLATE)
    fp2 = footer.add_paragraph()
    fp2.paragraph_format.tab_stops.add_tab_stop(Cm(right), WD_TAB_ALIGNMENT.RIGHT)
    r = fp2.add_run("%s\t" % B.BRANDS[b]["url"])
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = rgb(B.ORANGE_TEXT)
    r = fp2.add_run()
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(B.SLATE)
    field(r, "PAGE")
    if landscape:
        anchor_picture(fp2.add_run(), os.path.join(ART, "a4_slab_br.png"), 6.0, 23.7, 20.4, descr="", decorative=True)
    else:
        anchor_picture(fp2.add_run(), os.path.join(ART, "a4_slab_br.png"), 6.0, 15.0, 29.1, descr="", decorative=True)


def cover_header(b, header, photo_cover=True):
    hp = header.paragraphs[0]
    hp.text = ""
    r = hp.add_run()
    anchor_picture(r, os.path.join(ART, "a4_cover_%s.jpg" % b), 21.0, 0, 0, descr="", decorative=True, height_cm=29.7)
    r = hp.add_run()
    r.add_picture(B.logo(b, "reversed-dark", 1024), width=Cm(5.2))
    r._r.find(".//" + qn("wp:docPr")).set("descr", C.fill(WC.shared("alt.logo"), brand=b))


# ---------- documents ----------

def new_doc():
    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    page_setup(sec)
    doc.core_properties.language = "en-GB"
    return doc


def cover_front(doc, b, info, name):
    """The cover page text: label, title, subtitle and the meta line over the Dark Brown fade, then a page break."""
    ed = B.BRANDS[b]["edition"]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(318)
    r = p.add_run(info["label"])
    set_font(r, B.DISPLAY, 10, bold=True, colour=B.WHITE, caps=True, spacing=1.0)
    ppr = p._p.get_or_add_pPr()
    # sticker look: Climb Orange shading on the label run (White 10pt bold is too small for orange; use the run shading
    # on Dark Brown with an orange rule instead).
    border(p._p, "left", B.ORANGE, 48, 8)
    p = doc.add_paragraph(style="Title")
    rich(p, info["title"].replace("%s", name), base=B.WHITE)
    for r in p.runs:
        set_font(r, B.DISPLAY_BLACK, 38)
    content_control(p, "Document title", "")
    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run(info["sub"])
    r.font.color.rgb = rgb(B.GREY)
    content_control(p, "Subtitle", "")
    p = doc.add_paragraph()
    r = p.add_run(ed + ". " + B.BRANDS[b]["line"] + "." if ed else info["meta"])
    set_font(r, B.BODY, 10.5, bold=True, colour=B.TINT)
    if ed:
        p = doc.add_paragraph()
        r = p.add_run(info["meta"])
        set_font(r, B.BODY, 10, colour=B.GREY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def cover_doc(b, name):
    """The report and the proposal: a photo cover, a contents page and inner pages that show the styles in use."""
    words, t = original_words(b, name)
    row = [r for r in K.rows("word", b, original=True) if r["name"] == name][0]
    info = WC.cover(b, row)
    doc = new_doc()
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    cover_header(b, sec.first_page_header)
    standard_header(b, sec.header, sec.footer, label=info["label"])
    sec.first_page_footer.paragraphs[0].text = ""
    cover_front(doc, b, info, b)

    # Inner pages: a style guide in use.
    doc.add_paragraph(WC.shared("toc.eyebrow"), style="Eyebrow")
    toc = doc.add_paragraph()
    r = toc.add_run()
    field(r, 'TOC \\o "1-2" \\h \\z \\u')
    r.text = ""
    toc.add_run(WC.shared("toc.hint")).italic = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph(t("body.eyebrow"), style="Eyebrow")
    h = doc.add_paragraph(style="Heading 1")
    rich(h, C.rich(t("body.h1")), orange=B.ORANGE)
    doc.add_paragraph(t("body.lead"), style="Lead")
    doc.add_paragraph(t("body.p"))
    doc.add_paragraph(t("body.h2"), style="Heading 2")
    for item in C.items(words.one("body.bullets")):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(t("body.h3"), style="Heading 3")
    doc.add_paragraph(t("body.callout"), style="Callout")
    doc.add_paragraph(t("body.quote"), style="Quote")
    doc.add_paragraph(t("body.quote.source"), style="Quote Source")

    doc.add_paragraph(t("body.table-heading"), style="Heading 2")
    rows = C.table(words, words.one("body.table"))
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = doc.styles["Climb Table"]
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
    tblpr = tbl._tbl.tblPr
    look = tblpr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tblpr.append(look)
    for k, v in (("w:val", "04A0"), ("w:firstRow", "1"), ("w:lastRow", "0"), ("w:firstColumn", "0"), ("w:lastColumn", "0"),
                 ("w:noHBand", "0"), ("w:noVBand", "1")):
        look.set(qn(k), v)
    header_row(tbl)
    doc.add_paragraph(t("body.caption"), style="Caption")
    big = words.one("body.big")
    doc.add_paragraph(big.title, style="Big Number")
    doc.add_paragraph(big.text)
    doc.add_paragraph(t("body.next-heading"), style="Heading 2")
    doc.add_paragraph(t("body.next"))
    return doc


def original_words(b, name):
    """The copy file of one of the first set, and a reader for one of its sections."""
    row = [r for r in K.rows("word", b, original=True) if r["name"] == name][0]
    words = WC.copy_doc(b, row)
    return words, lambda path: words.one(path).text


def letterhead(b):
    words, t = original_words(b, "Letterhead")
    doc = new_doc()
    sec = doc.sections[0]
    standard_header(b, sec.header, sec.footer)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    for line in t("letter.recipient").split("\n"):
        p.add_run(line).add_break()
    content_control(p, "Recipient", "")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run(t("letter.date"))
    content_control(p, "Date", "")
    p = doc.add_paragraph(t("letter.greeting"))
    p.paragraph_format.space_before = Pt(18)
    h = doc.add_paragraph(style="Heading 2")
    h.add_run(t("letter.subject"))
    for para in words.all("letter.p"):
        doc.add_paragraph(para.text)
    doc.add_paragraph(t("letter.sign-off"))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run(t("letter.name"))
    set_font(r, B.DISPLAY, 10.5, bold=True, colour=B.CHARCOAL)
    r.add_break()
    r = p.add_run(t("letter.role"))
    r.font.color.rgb = rgb(B.SLATE)
    return doc


def press_release(b):
    words, t = original_words(b, "Press-release")
    doc = new_doc()
    sec = doc.sections[0]
    standard_header(b, sec.header, sec.footer, label=t("header.label"))
    doc.add_paragraph(t("body.eyebrow"), style="Eyebrow")
    title = doc.add_paragraph(style="Title")
    rich(title, C.rich(t("body.title")))
    for r in title.runs:
        set_font(r, B.DISPLAY_BLACK, 28)
    content_control(title, "Headline", "")
    doc.add_paragraph(t("body.lead"), style="Lead")
    for para in words.all("body.p"):
        doc.add_paragraph(para.text)
    doc.add_paragraph(t("body.quote"), style="Quote")
    doc.add_paragraph(t("body.quote.source"), style="Quote Source")
    doc.add_paragraph(t("body.close"))
    p = doc.add_paragraph(t("body.ends"), style="Eyebrow")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(t("notes.heading"), style="Heading 2")
    for item in C.items(words.one("notes.bullets")):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(t("contact.heading"), style="Heading 3")
    doc.add_paragraph(t("contact.text"))
    return doc


def meeting_notes(b):
    words, t = original_words(b, "Meeting-notes")
    doc = new_doc()
    sec = doc.sections[0]
    standard_header(b, sec.header, sec.footer, label=t("header.label"))
    doc.add_paragraph(t("body.eyebrow"), style="Eyebrow")
    title = doc.add_paragraph(style="Title")
    rich(title, C.rich(t("body.title")))
    for r in title.runs:
        set_font(r, B.DISPLAY_BLACK, 26)
    content_control(title, "Meeting name", "")
    kv = [line.split(": ", 1) for line in C.items(words.one("body.kv"))]
    tbl = doc.add_table(rows=len(kv), cols=2)
    tbl.style = doc.styles["Climb Table"]
    for i, (k, v) in enumerate(kv):
        tbl.cell(i, 0).text = k
        tbl.cell(i, 1).text = v
    tblpr = tbl._tbl.tblPr
    look = OxmlElement("w:tblLook")
    for k, v in (("w:val", "0020"), ("w:firstRow", "0"), ("w:firstColumn", "1"), ("w:noHBand", "0"), ("w:noVBand", "1")):
        look.set(qn(k), v)
    tblpr.append(look)
    doc.add_paragraph(t("agenda.heading"), style="Heading 2")
    for item in C.items(words.one("agenda.items")):
        doc.add_paragraph(item, style="List Number")
    doc.add_paragraph(t("decisions.heading"), style="Heading 2")
    for item in C.items(words.one("decisions.items")):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(t("actions.heading"), style="Heading 2")
    rows = C.table(words, words.one("actions.table"))
    a = doc.add_table(rows=len(rows), cols=len(rows[0]))
    a.style = doc.styles["Climb Table"]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            a.cell(i, j).text = v
    header_row(a)
    doc.add_paragraph(t("next.heading"), style="Heading 2")
    doc.add_paragraph(t("next.text"))
    return doc


ORIGINAL = {
    "Letterhead": letterhead,
    "Report": lambda b: cover_doc(b, "Report"),
    "Proposal": lambda b: cover_doc(b, "Proposal"),
    "Press-release": press_release,
    "Meeting-notes": meeting_notes,
}


# ---------- the catalogue documents, rendered from word_content.py ----------

def table_look(tbl, first_row=True, first_col=False):
    tblpr = tbl._tbl.tblPr
    look = tblpr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tblpr.append(look)
    for k, v in (("w:val", "04A0" if first_row else "0020"), ("w:firstRow", "1" if first_row else "0"), ("w:lastRow", "0"),
                 ("w:firstColumn", "1" if first_col else "0"), ("w:lastColumn", "0"), ("w:noHBand", "0"), ("w:noVBand", "1")):
        look.set(qn(k), v)


def fill_table(doc, header, rows, widths=None, first_col=False):
    """A Climb Table with a repeating header row; narrow type when the table is wide."""
    cols = len(header) if header else len(rows[0])
    tbl = doc.add_table(rows=len(rows) + (1 if header else 0), cols=cols)
    tbl.style = doc.styles["Climb Table"]
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    size = 9.5 if cols < 7 else (8 if cols < 10 else 7.5)
    all_rows = ([header] if header else []) + rows
    for i, row in enumerate(all_rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            rich(p, str(v), size=size)
            if first_col and j == 0:
                for r in p.runs:
                    set_font(r, B.DISPLAY, size, bold=True, colour=B.CHARCOAL)
            if widths:
                cell.width = Cm(widths[j])
    if widths:
        for j, w in enumerate(widths):
            tbl.columns[j].width = Cm(w)
    if header:
        header_row(tbl)
    if cols >= 7:
        for row in tbl.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    table_look(tbl, first_row=bool(header), first_col=first_col)
    return tbl


def merge_field(run, name):
    """A mail merge field with its placeholder result, so the slot shows before the merge. Each part is its own run
    carrying the run's formatting, which is what Word and LibreOffice both apply to the result."""
    import copy
    rpr = run._r.find(qn("w:rPr"))
    parent = run._r.getparent()
    at = list(parent).index(run._r)

    def run_el():
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        return r
    parts = []
    for kind, text in (("begin", None), ("instr", 'MERGEFIELD %s \\* MERGEFORMAT' % name), ("separate", None),
                       ("text", "\u00ab%s\u00bb" % name), ("end", None)):
        r = run_el()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " %s " % text
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r.append(el)
        parts.append(r)
    parent.remove(run._r)
    for k, r in enumerate(parts):
        parent.insert(at + k, r)


def merge_text(p, text, size=None, colour=None):
    """Text with «Field» tokens turned into mail merge fields."""
    import re
    for tok in re.split(r"(«[^»]+»)", text):
        if not tok:
            continue
        r = p.add_run()
        if tok.startswith("«"):
            merge_field(r, tok[1:-1])
        else:
            r.text = tok
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = rgb(colour)


def restart_numbering(doc, paragraphs):
    """Give a numbered list its own numbering instance, so it starts at 1 instead of continuing the last list."""
    numbering = doc.part.numbering_part.element
    style_pr = doc.styles["List Number"].element.pPr
    num_id = style_pr.numPr.numId.val
    abstract = numbering.num_having_numId(num_id).abstractNumId.val
    new = numbering.add_num(abstract)
    new.add_lvlOverride(0).add_startOverride(1)
    for p in paragraphs:
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = new.numId


def sign_block(doc, parties):
    tbl = doc.add_table(rows=1, cols=len(parties))
    tbl.style = doc.styles["Climb Table"]
    tbl.autofit = False
    for j, party in enumerate(parties):
        cell = tbl.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(C.fill(WC.shared("sign.party"), party=party))
        set_font(r, B.DISPLAY, 9.5, bold=True, colour=B.CHARCOAL)
        for line in C.items(C.load(WC.SHARED).one("sign.lines")):
            q = cell.add_paragraph()
            q.paragraph_format.space_before = Pt(14)
            r = q.add_run("%s:  " % line)
            set_font(r, B.BODY, 9, bold=True, colour=B.SLATE)
            r = q.add_run("_" * 28)
            r.font.color.rgb = rgb(B.SLATE)
        cell.width = Cm(16.6 / len(parties))
    table_look(tbl, first_row=False)


def cert_page(doc, b, eyebrow, line, after, fields, signer):
    """One certificate: the mark, the eyebrow, the name as merge fields, the line, a signature and the lockup."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run()
    r.add_picture(B.logo(b, "full-colour-light", 1024), width=Cm(4.6))
    r._r.find(".//" + qn("wp:docPr")).set("descr", C.fill(WC.shared("alt.logo"), brand=b))
    p = doc.add_paragraph(style="Eyebrow")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.add_run(eyebrow)
    p = doc.add_paragraph(style="Lead")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    for k, (fld, _) in enumerate(fields):
        if k:
            p.add_run(" ")
        r = p.add_run()
        set_font(r, B.DISPLAY_BLACK, 36, colour=B.CHARCOAL, caps=True, spacing=-0.4)
        merge_field(r, fld)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(os.path.join(ART, "rule_short.png"), width=Cm(2.6))
    r._r.find(".//" + qn("wp:docPr")).set("descr", "")
    p = doc.add_paragraph(style="Lead")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    merge_text(p, after)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run("_" * 34)
    r.font.color.rgb = rgb(B.SLATE)
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(signer)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run()
    r.add_picture(B.logo(b, "tagline-lockup", 1024), width=Cm(5.0))
    r._r.find(".//" + qn("wp:docPr")).set("descr", C.fill(WC.shared("alt.lockup"), brand=b, line=B.BRANDS[b]["line"]))


def render_blocks(doc, b, blocks):
    for blk in blocks:
        kind = blk[0]
        if kind == "eyebrow":
            doc.add_paragraph(blk[1], style="Eyebrow")
        elif kind == "title":
            p = doc.add_paragraph(style="Title")
            rich(p, blk[1])
            for r in p.runs:
                set_font(r, B.DISPLAY_BLACK, 26)
        elif kind == "sub":
            doc.add_paragraph(blk[1], style="Subtitle")
        elif kind == "lead":
            rich(doc.add_paragraph(style="Lead"), blk[1])
        elif kind in ("h1", "h2", "h3"):
            rich(doc.add_paragraph(style="Heading %s" % kind[1]), blk[1])
        elif kind == "p":
            for line in blk[1].split("\n"):
                rich(doc.add_paragraph(), line)
        elif kind == "bullets":
            for t in blk[1]:
                rich(doc.add_paragraph(style="List Bullet"), t)
        elif kind == "numbers":
            restart_numbering(doc, [rich(doc.add_paragraph(style="List Number"), t) for t in blk[1]])
        elif kind == "checklist":
            for t in blk[1]:
                doc.add_paragraph("[  ]  " + t)
        elif kind == "callout":
            rich(doc.add_paragraph(style="Callout"), blk[1])
        elif kind == "quote":
            doc.add_paragraph(blk[1], style="Quote")
            doc.add_paragraph(blk[2], style="Quote Source")
        elif kind == "big":
            doc.add_paragraph(blk[1], style="Big Number")
            doc.add_paragraph(blk[2])
        elif kind == "small":
            doc.add_paragraph(blk[1], style="Small Print")
        elif kind == "kv":
            fill_table(doc, None, [[k, v] for k, v in blk[1]], [4.6, 12.0], first_col=True)
        elif kind == "table":
            fill_table(doc, blk[1], blk[2], blk[3] if len(blk) > 3 else None)
        elif kind == "clauses":
            n = blk[2] if len(blk) > 2 else 1
            for k, (heading, note) in enumerate(blk[1], n):
                doc.add_paragraph("%d. %s" % (k, heading), style="Heading 2")
                doc.add_paragraph(WC.shared("clause.text"))
                doc.add_paragraph(WC.shared("clause.hint") + " " + note, style="Caption")
        elif kind == "sign":
            doc.add_paragraph(WC.shared("sign.heading"), style="Heading 2")
            sign_block(doc, blk[1])
        elif kind == "sigline":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(24)
            signed, dated = WC.shared("sigline.signed"), WC.shared("sigline.date")
            r = p.add_run(signed + "  " + "_" * 30 + "      " + dated + "  " + "_" * 14)
            r.font.color.rgb = rgb(B.SLATE)
            doc.add_paragraph(blk[1], style="Caption")
        elif kind == "page":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "toc":
            doc.add_paragraph(WC.shared("toc.eyebrow"), style="Eyebrow")
            toc = doc.add_paragraph()
            r = toc.add_run()
            field(r, 'TOC \\o "1-2" \\h \\z \\u')
            r.text = ""
            toc.add_run(WC.shared("toc.hint")).italic = True
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "cert":
            cert_page(doc, b, *blk[1:])
        else:
            raise ValueError("unknown block %r" % (kind,))


def landscape(sec, top=2.0, bottom=1.8):
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(top), Cm(bottom)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)


def catalogue_doc(b, row):
    kind = WC.DOCS[row["name"]]
    doc = new_doc()
    sec = doc.sections[0]
    if kind == "cover":
        info = WC.cover(b, row)
        sec.different_first_page_header_footer = True
        cover_header(b, sec.first_page_header)
        standard_header(b, sec.header, sec.footer, label=info["label"])
        sec.first_page_footer.paragraphs[0].text = ""
        cover_front(doc, b, info, b)
    elif kind == "plain":
        standard_header(b, sec.header, sec.footer, label=WC.label(b, row))
    elif kind == "landscape-plain":
        landscape(sec, top=3.0, bottom=2.4)
        standard_header(b, sec.header, sec.footer, label=WC.label(b, row), landscape=True)
    elif kind == "landscape":
        landscape(sec)
        hp = sec.header.paragraphs[0]
        hp.text = ""
        anchor_picture(hp.add_run(), os.path.join(ART, "a4_slab_tl.png"), 3.6, 0, 0, descr="", decorative=True)
        anchor_picture(hp.add_run(), os.path.join(ART, "a4_slab_br.png"), 6.0, 23.7, 20.4, descr="", decorative=True)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        r = fp.add_run(B.LEGAL)
        r.font.size = Pt(6.5)
        r.font.color.rgb = rgb(B.SLATE)
    render_blocks(doc, b, WC.blocks(b, row))
    return doc


def save(doc, b, row):
    stem = K.stem(b, row)
    d = os.path.join(B.OUT, "office", K.folder(b, row))
    os.makedirs(d, exist_ok=True)
    doc.core_properties.title = "%s %s" % (b, row["title"].lower())
    doc.core_properties.author = b
    p = os.path.join(d, stem + ".dotx")
    doc.save(p)
    apply_theme(p, b)
    os.makedirs(os.path.join(B.OUT, "_preview"), exist_ok=True)
    to_document(p, os.path.join(B.OUT, "_preview", stem + ".docx"))
    print(p)


def main():
    args = sys.argv[1:]
    if args[:1] == ["--copy"]:
        print("\n".join(WC.all_copy(original=True)))
        return
    only = None
    if args[:1] == ["--only"]:
        only, args = args[1:], []
    brands = args or list(B.BRANDS)
    for b in brands:
        for row in K.rows("word", b):
            if only and row["name"] not in only:
                continue
            fn = ORIGINAL.get(row["name"]) if row["original"] else None
            doc = fn(b) if fn else catalogue_doc(b, row)
            save(doc, b, row)


if __name__ == "__main__":
    main()
