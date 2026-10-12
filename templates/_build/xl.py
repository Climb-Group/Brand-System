"""Shared pieces of the Excel templates: fonts, fills, the banner, the table sheet, the summary chart and the template
packaging. excel_build.py (the first set) and excel_books.py (the catalogue) build on these."""
import os
import zipfile

from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import brand as B
import ooxml_theme as T

ART = os.path.join(B.OUT, "_art")
GBP = '"\u00a3"#,##0;-"\u00a3"#,##0'
GBP_PLAIN = '"\u00a3"#,##0'

HEAD_FILL = PatternFill("solid", fgColor=B.DARK)
HEAD_FONT = Font(name=B.DISPLAY, size=9, bold=True, color=B.WHITE)
BODY_FONT = Font(name=B.BODY, size=10, color=B.SLATE)
TITLE_FONT = Font(name=B.DISPLAY_XB, size=20, color=B.CHARCOAL)
SUB_FONT = Font(name=B.BODY, size=10.5, color=B.SLATE)
EYEBROW_FONT = Font(name=B.BODY, size=9, bold=True, color=B.ORANGE_TEXT)
TOTAL_FONT = Font(name=B.DISPLAY, size=10, bold=True, color=B.CHARCOAL)
THIN = Side(style="thin", color=B.GREY)
TOTAL_TOP = Side(style="medium", color=B.ORANGE)


def banner(ws, b, cols):
    img = XLImage(os.path.join(ART, "xl_banner_%s.png" % b))
    width_px = sum((ws.column_dimensions[get_column_letter(c)].width or 10) * 7.2 for c in range(1, cols + 1))
    img.width = int(width_px)
    img.height = int(width_px / 10)
    ws.add_image(img, "A1")
    ws.row_dimensions[1].height = img.height * 0.75
    return img


def sheet(wb, b, name, eyebrow, title, sub, columns, rows, validations=(), widths=None, totals=None, number_cols=(),
          date_cols=(), time_cols=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = B.ORANGE
    for i, w in enumerate(widths or [18] * len(columns), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    banner(ws, b, len(columns))
    ws["A3"], ws["A4"], ws["A5"] = eyebrow.upper(), title.upper(), sub
    ws["A3"].font, ws["A4"].font, ws["A5"].font = EYEBROW_FONT, TITLE_FONT, SUB_FONT
    ws.row_dimensions[4].height = 30
    hr = 7
    for j, c in enumerate(columns, 1):
        cell = ws.cell(hr, j, c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[hr].height = 26
    for i, row in enumerate(rows, hr + 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if j in number_cols:
                cell.number_format = GBP
            if j in date_cols:
                cell.number_format = "ddd d mmm yyyy"
            if j in time_cols:
                cell.number_format = "hh:mm"
    last = hr + len(rows)
    ref = "A%d:%s%d" % (hr, get_column_letter(len(columns)), last)
    t = Table(displayName=name.replace(" ", ""), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True, showColumnStripes=False)
    ws.add_table(t)
    for col_letter, options, prompt in validations:
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True,
                            promptTitle="Pick one", prompt=prompt, showInputMessage=True, showErrorMessage=True,
                            errorTitle="Not on the list", error="Pick a value from the list.")
        dv.add("%s%d:%s%d" % (col_letter, hr + 1, col_letter, last + 200))
        ws.add_data_validation(dv)
    if totals:
        tr = last + 2
        ws.cell(tr, 1, "Total").font = TOTAL_FONT
        for col in totals:
            cl = get_column_letter(col)
            c = ws.cell(tr, col, "=SUBTOTAL(109,%s[%s])" % (t.displayName, columns[col - 1]))
            c.font, c.number_format = TOTAL_FONT, GBP
        for col in range(1, len(columns) + 1):
            ws.cell(tr, col).border = Border(top=TOTAL_TOP)
    ws.freeze_panes = "A%d" % (hr + 1)
    ws.print_title_rows = "%d:%d" % (hr, hr)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = B.BRANDS[b]["url"]
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.left.font = ws.oddFooter.right.font = B.BODY
    return ws, t, hr, last


def summary_chart(ws, title, cats_ref, vals_ref, anchor, y_title="\u00a3"):
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = title
    ch.y_axis.title = y_title
    ch.legend = None
    ch.add_data(vals_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    s = ch.series[0]
    s.graphicalProperties.solidFill = B.ORANGE
    s.graphicalProperties.line.solidFill = B.ORANGE
    ch.width, ch.height = 18, 9
    ws.add_chart(ch, anchor)


def finish(path, b, template=True):
    """Swap in the brand theme and mark the package as a template."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/theme/theme1.xml":
                data = T.theme_xml(b).encode()
            elif item.filename == "[Content_Types].xml" and template:
                data = data.replace(b"spreadsheetml.sheet.main+xml", b"spreadsheetml.template.main+xml")
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"spreadsheetml.template.main+xml", b"spreadsheetml.sheet.main+xml")
            zout.writestr(item, data)
    os.replace(tmp, path)
